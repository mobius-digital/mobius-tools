import openpyxl, re, json, sys, os
HERE = os.path.dirname(os.path.abspath(__file__))
MONTHS = ['January','February','March','April','May','June','July','August','September','October','November','December']
TAX2BUCKET = {'Software & subscriptions':'Software','Contract labor (1099)':'Contractors','Advertising & marketing':'Ads/Marketing',
  'Bank & merchant fees':'Merchant fee','Meals':'Meals & entertainment','Office supplies & equipment':'Office & equipment',
  'Product testing':'Product testing','Travel & gas':'Travel','Utilities & phone':'Utilities & phone','Dues & memberships':'Dues & memberships',
  'Taxes & licenses':'Taxes & licenses','Personal':'Personal (not a business cost)','Client revenue':'Revenue','Other':'Other'}

def num(v):
    if v is None or v == '': return None
    if isinstance(v, (int, float)): return float(v)
    s = str(v).replace('$','').replace(',','').strip()
    try: return float(s)
    except: return None

# ---- revenue: which client a line belongs to ----
CLIENTS = [
  (r'lucky golf|luck golf|\bLG\b', 'Lucky Golf'), (r'party pa+tch|part patch|\bPP\b', 'Party Patch'),
  (r'grunk', 'Grunk Dolfer'), (r'dartee', 'Dartee Golf'), (r'bonk|floyd', 'Bonk Golf'),
  (r'instant ?iv', 'InstantIV'), (r'camskns', 'CAMSKNS'), (r'sole ami', 'Sole Amié'),
  (r'booster theme|boostr theme', 'Booster Theme'), (r'speed ?boostr', 'Speed Boostr'), (r'tork ?strap', 'TorkStrap'),
  (r'lasting portrait', 'Lasting Portrait'), (r'burke golf|mike', 'Burke Golf'), (r'brew ?craft', 'BrewCraft'),
  (r'sunshine sauce', 'Sunshine Sauce'), (r'judyp', 'JudyP Apparel'), (r'le pickle', 'Le Pickle Club'),
  (r'ear-goes|ear goes', 'Ear-Goes'), (r'bryce', 'Bryce (consulting)'), (r'galway', 'Galway Bay Apparel'),
  (r'golf sock', 'The Golf Sock'), (r'united patriot', 'United Patriots'), (r'vetripaws', 'VetriPaws'), (r'speedin', 'SpeedIn'),
]
def client_of(name):
    for pat, c in CLIENTS:
        if re.search(pat, name, re.I): return c
    return None

PEOPLE = [(r'ahsan', 'Syed Ahsan Raza'), (r'raphael', 'Raphael'), (r'ravo', 'Ravo Schubert'), (r'srully', 'Srully'),
  (r'william', 'William'), (r'melanie', 'Melanie'), (r'cooper', 'Cooper'), (r'noma', 'Noma'), (r'radhesh', 'Radhesh Gowd')]
ROLES = [(r'creative strat', 'Creative strategist'), (r'media buy', 'Media buyer'), (r'graphic', 'Graphic designer'),
  (r'video edit', 'Video editor'), (r'\bva\b', 'Outbound VA'), (r'consultant', 'Consultant'), (r'website', 'Website contractor')]
def person_of(name):
    for pat, c in PEOPLE:
        if re.search(pat, name, re.I): return c
    for pat, c in ROLES:
        if re.search(pat, name, re.I): return c + ' (unnamed)'
    return name

# ---- the "Other" column: pick the CPA line from the words ----
OTHER = [
  (r'charity|^cpa$|^\(no label\)$|book', 'Other'),
  (r'haircut', 'Advertising & marketing'),
  (r'laptop|desk|monitor|camera|recording|microphone|tripod|lamp|equipment|office|supplies|supply|blue light|amazon|chair|ipad|keyboard|dock|neon sign|co2|speakers', 'Office supplies & equipment'),
  (r'dinner|lunch|breakfast|food|meal|coffee|restaurant|creator conversation', 'Meals'),
  (r'hotel|uber|lyft|flight|airfare|business trip|gas\b|travel|parking', 'Travel & gas'),
  (r'llc|registered agent|tax fee|license|filing', 'Taxes & licenses'),
  (r'membership|amex gold', 'Dues & memberships'),
  (r'phone|internet|wifi|t-mobile|verizon', 'Utilities & phone'),
  (r'product testing|simplisoda', 'Product testing'),
  (r'golf content|video content|content|client gift|gifts|brand marketing|shipping costs|entertainment|golf\b|replo|landing page', 'Advertising & marketing'),
  (r'consultant|freelanc|va\b', 'Contract labor (1099)'),
  (r'seamless|online jobs|software', 'Software & subscriptions'),
  (r'course|agency lab|coaching|education', 'Other'),
]
def other_cat(name):
    for pat, c in OTHER:
        if re.search(pat, name, re.I): return c
    return 'Other'

def parse_year(year, path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out, checks = [], []
    for mi, mname in enumerate(MONTHS):
        if mname not in wb.sheetnames: continue
        ws = wb[mname]
        month = f'{year}-{mi+1:02d}'
        rows = list(ws.iter_rows(values_only=True))
        summ = {}
        for r in rows:
            if len(r) > 14 and isinstance(r[13], str):
                v = num(r[14])
                if v is not None: summ.setdefault(r[13].strip(), v)
        items = []
        for r in rows[2:]:
            r = list(r) + [None] * 14
            for col, kind in ((0,'rev'),(2,'software'),(4,'contractor'),(6,'payroll'),(8,'ads'),(10,'other')):
                name, amt = r[col], num(r[col+1])
                if amt is None or abs(amt) < 0.005: continue
                name = (str(name).strip() if name not in (None, '') else '(no label)')
                items.append((kind, name, round(amt, 2)))
        rev = round(sum(a for k,_,a in items if k=='rev'), 2)
        exp = round(sum(a for k,_,a in items if k!='rev'), 2)
        fee = next((v for k,v in summ.items() if k.startswith('Merchant Fee')), 0.0) or 0.0
        s_rev, s_exp = summ.get('Revenue'), summ.get('Total Expenses')
        checks.append((month, rev, s_rev, exp, s_exp, round(fee,2)))
        for kind, name, amt in items:
            note = f'{year} Master Finance sheet: "{name}"'
            if kind == 'rev':
                low = name.lower()
                if 'refund' in low and not client_of(name.replace('Refund','')) or re.search(r'amazon refund|agency lab refund', low):
                    # money back on a purchase is a smaller expense, not revenue
                    vendor = re.sub(r'\s*refunds?\s*', ' ', name, flags=re.I).strip(' |') or name
                    tax = {'amazon': 'Office supplies & equipment', 'replo': 'Advertising & marketing'}.get(vendor.lower(), 'Other')
                    out.append(dict(month=month, type='out', vendor=vendor, amount=-amt, tax_cat=tax, note=note + ' (refund received)'))
                    continue
                c = client_of(name)
                if c is None: c = name.split('|')[0].strip() or name
                out.append(dict(month=month, type='in', vendor=c, amount=amt, tax_cat='Client revenue', note=note if name != c else None))
            elif kind == 'software':
                out.append(dict(month=month, type='out', vendor=name, amount=amt, tax_cat='Software & subscriptions', note=None))
            elif kind in ('contractor','payroll'):
                who = person_of(name)
                out.append(dict(month=month, type='out', vendor=who, amount=amt, tax_cat='Contract labor (1099)', note=note if who != name else None))
            elif kind == 'ads':
                out.append(dict(month=month, type='out', vendor=name, amount=amt, tax_cat='Advertising & marketing', note=None))
            else:
                cat = other_cat(name)
                who = person_of(name) if cat == 'Contract labor (1099)' else name
                out.append(dict(month=month, type='out', vendor=who, amount=amt, tax_cat=cat, note=note if who != name else None))
        if fee:
            out.append(dict(month=month, type='fee', vendor='Stripe', amount=round(fee,2), tax_cat='Bank & merchant fees',
                            note=f'{year} Master Finance sheet merchant fee' + (' (estimated at 2.93%)' if year < 2025 else '')))
    return out, checks

if __name__ == '__main__':
    allrows, bad = [], []
    for y in (2023, 2024, 2025):
        rows, checks = parse_year(y, os.path.join(HERE, f'{y}.xlsx'))
        allrows += rows
        for m, rev, srev, exp, sexp, fee in checks:
            ok = (srev is None or abs(rev - srev) < 0.02) and (sexp is None or abs(exp - sexp) < 0.02)
            print(f'{m}  rev {rev:>10.2f} sheet {srev!s:>10}  exp {exp:>10.2f} sheet {sexp!s:>10}  fee {fee:>8.2f}  {"OK" if ok else "MISMATCH"}')
            if not ok: bad.append(m)
    json.dump(allrows, open(os.path.join(HERE, 'rows.json'), 'w'), indent=0)
    print('rows', len(allrows), 'mismatch', bad)
